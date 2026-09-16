import os
import sys
import re
import time
import sqlite3
import threading
from typing import List, Dict, Any, Optional, Callable
import pandas as pd

# Add root directory to sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

from backend.connectors.access_connector import AccessConnector
from backend.connectors.excel_connector import ExcelConnector

try:
    import openpyxl
except ImportError:
    openpyxl = None

try:
    import xlrd
except ImportError:
    xlrd = None

CONSOLIDATED_DB_PATH = os.path.join(BASE_DIR, "data", "consolidated_lands.db")

# Global sync state tracker
_SYNC_STATE: Dict[str, Any] = {
    "status": "idle",  # "idle", "running", "completed", "error"
    "current_source": "",
    "current_table": "",
    "tables_completed": 0,
    "total_tables": 0,
    "rows_copied": 0,
    "start_time": None,
    "elapsed_seconds": 0,
    "error_message": None,
    "summary": []
}
_SYNC_LOCK = threading.Lock()

def get_sync_status() -> Dict[str, Any]:
    with _SYNC_LOCK:
        state = dict(_SYNC_STATE)
        if state["status"] == "running" and state["start_time"]:
            state["elapsed_seconds"] = round(time.time() - state["start_time"], 1)
        return state

def _update_sync_state(**kwargs):
    with _SYNC_LOCK:
        _SYNC_STATE.update(kwargs)
        if _SYNC_STATE["start_time"]:
            _SYNC_STATE["elapsed_seconds"] = round(time.time() - _SYNC_STATE["start_time"], 1)

def sanitize_identifier(name: str) -> str:
    """Converts raw table or column names to clean, valid SQLite identifiers"""
    clean = re.sub(r'[^a-zA-Z0-9_]', '_', str(name).strip())
    clean = re.sub(r'_+', '_', clean).strip('_')
    if not clean or clean[0].isdigit():
        clean = f"col_{clean}"
    return clean.lower()

class ETLPipeline:
    """Extracts tables from network Access and Excel sources into a local consolidated SQLite DB"""

    @staticmethod
    def get_database_path() -> str:
        return CONSOLIDATED_DB_PATH

    @staticmethod
    def run_full_sync(on_progress: Optional[Callable[[Dict[str, Any]], None]] = None) -> Dict[str, Any]:
        """Runs complete extraction and consolidation across all sources"""
        os.makedirs(os.path.dirname(CONSOLIDATED_DB_PATH), exist_ok=True)
        start_time = time.time()

        _update_sync_state(
            status="running",
            current_source="Initializing",
            current_table="",
            tables_completed=0,
            total_tables=0,
            rows_copied=0,
            start_time=start_time,
            elapsed_seconds=0,
            error_message=None,
            summary=[]
        )

        sources_to_sync = [
            {
                "name": "Offer Letters 2026 (Live System)",
                "type": "access",
                "path": r"\\Lands02\lims\OfferLetters 2026.accdb",
                "password": None,
                "prefix": "ol2026",
                "tables": [
                    "Beneficiary_Details",
                    "A1_Subdivision_Register",
                    "Farm_register",
                    "PermitsPrinted",
                    "Allocation_Register",
                    "Planned_Farm",
                    "Company",
                    "Withdrawals",
                    "Intention"
                ]
            },
            {
                "name": "LANDS 02 Network Database (LIMS Live Old)",
                "type": "access",
                "path": r"\\Lands02\lims\backups\lims live\old system\LANDS 02_Backup.mdb",
                "password": "ALLOCATIONS",
                "prefix": "lands02",
                "tables": [
                    "FARM DETAILS",
                    "PERSONAL DETAILS",
                    "Farm Withdrawn"
                ]
            },
            {
                "name": "A2 Allocations Database (Mrs Gavi Excel)",
                "type": "excel",
                "path": r"\\Lands02\offer letters\MRS GAVI\Copy of A2 ALLOC-2 Old Database (version 1) (Autosaved)-3.xls",
                "password": None,
                "prefix": "a2_alloc",
                "tables": [
                    "Sheet1",
                    "Sheet2"
                ]
            },
            {
                "name": "Current Farm Registers as of 2025",
                "type": "excel_folder",
                "path": r"\\Lands02\lims\Zec\Current Farm Registers as of 2025",
                "cache_path": os.path.join(BASE_DIR, "data", "network_cache", "farm_registers_2025"),
                "password": None,
                "prefix": "reg2025",
                "workbooks": [
                    ("manicaland", "Manicaland Farm register.xlsx"),
                    ("masheast", "MASH EAST Copy of revised_farm_register_oct_2024(1) - Copy (1).xlsx"),
                    ("mashcentral", "Mashonaland Central Province Updated Registers 2024 (1).xlsx"),
                    ("mashwest", "Mashonaland  West Farm Registers 2024.xlsx"),
                    ("masvingo", "Masvingo Farm Register Mrs Muchemwa date 18Feb2025 (1).xls"),
                    ("matnorth", "Mat.North Farm Registers January 2025.xls"),
                    ("matsouth", "Matebeleland South Database 06_02_25 (1).xls"),
                    ("midlands", "Midlands_Farm Registers Updated (2024).xlsx")
                ]
            }
        ]

        total_tables = sum(len(s.get("tables", [])) or len(s.get("workbooks", [])) * 10 for s in sources_to_sync)
        _update_sync_state(total_tables=total_tables)

        # Connect to local SQLite
        sqlite_conn = sqlite3.connect(CONSOLIDATED_DB_PATH)
        sqlite_conn.execute("PRAGMA journal_mode = WAL;")
        sqlite_conn.execute("PRAGMA synchronous = NORMAL;")

        fts_entries = []
        summary = []
        total_rows_copied = 0
        tables_done = 0

        try:
            for src in sources_to_sync:
                src_name = src["name"]
                src_type = src["type"]
                file_path = src["path"]
                pwd = src.get("password")
                prefix = src["prefix"]
                tables = src["tables"]

                _update_sync_state(current_source=src_name)

                if not os.path.exists(file_path):
                    print(f"[ETL Warning] Source path unreachable: {file_path}")
                    summary.append({"source": src_name, "status": "UNREACHABLE", "rows": 0})
                    continue

                if src_type == "access":
                    conn, _ = AccessConnector._open_connection(file_path, password=pwd, timeout=20)
                    with conn:
                        cursor = conn.cursor()
                        for tbl in tables:
                            sqlite_tbl_name = f"{prefix}_{sanitize_identifier(tbl)}"
                            _update_sync_state(current_table=f"{src_name} -> {tbl}")
                            t0_tbl = time.time()

                            try:
                                rows_copied, sample_fts = ETLPipeline._sync_access_table(
                                    cursor=cursor,
                                    access_tbl=tbl,
                                    sqlite_conn=sqlite_conn,
                                    sqlite_tbl=sqlite_tbl_name,
                                    src_name=src_name
                                )
                                total_rows_copied += rows_copied
                                tables_done += 1
                                fts_entries.extend(sample_fts)
                                _update_sync_state(
                                    tables_completed=tables_done,
                                    rows_copied=total_rows_copied
                                )
                                elapsed_tbl = time.time() - t0_tbl
                                print(f"[ETL] Extracted {tbl} -> {sqlite_tbl_name}: {rows_copied:,} rows in {elapsed_tbl:.2f}s")
                                summary.append({
                                    "source": src_name,
                                    "original_table": tbl,
                                    "sqlite_table": sqlite_tbl_name,
                                    "rows": rows_copied,
                                    "status": "SUCCESS"
                                })
                            except Exception as tbl_err:
                                print(f"[ETL Error] Failed table {tbl}: {tbl_err}")
                                summary.append({
                                    "source": src_name,
                                    "original_table": tbl,
                                    "sqlite_table": sqlite_tbl_name,
                                    "rows": 0,
                                    "status": f"ERROR: {str(tbl_err)[:80]}"
                                })

                elif src_type == "excel_folder":
                    workbooks = src.get("workbooks", [])
                    cache_dir = src.get("cache_path")
                    net_dir = src.get("path")
                    
                    if cache_dir:
                        os.makedirs(cache_dir, exist_ok=True)
                        
                    for prov_key, fname in workbooks:
                        target_file = None
                        cache_file = os.path.join(cache_dir, fname) if cache_dir else None
                        net_file = os.path.join(net_dir, fname) if net_dir else None
                        
                        if cache_file and os.path.exists(cache_file):
                            target_file = cache_file
                        elif net_file and os.path.exists(net_file):
                            target_file = net_file
                            
                        if not target_file:
                            print(f"[ETL Warning] Workbook not found: {fname}")
                            continue
                            
                        _update_sync_state(current_table=f"{src_name} -> {fname}")
                        t0_wb = time.time()
                        try:
                            prov_rows, prov_fts, prov_tbls = ETLPipeline._sync_register_workbook(
                                file_path=target_file,
                                prov_key=prov_key,
                                sqlite_conn=sqlite_conn,
                                src_name=src_name
                            )
                            total_rows_copied += prov_rows
                            tables_done += prov_tbls
                            fts_entries.extend(prov_fts)
                            _update_sync_state(
                                tables_completed=tables_done,
                                rows_copied=total_rows_copied
                            )
                            elapsed_wb = time.time() - t0_wb
                            print(f"[ETL] Extracted {fname}: {prov_rows:,} rows across {prov_tbls} tables in {elapsed_wb:.2f}s")
                            summary.append({
                                "source": src_name,
                                "original_table": fname,
                                "sqlite_table": f"reg2025_{prov_key}_*",
                                "rows": prov_rows,
                                "status": "SUCCESS"
                            })
                        except Exception as wb_err:
                            print(f"[ETL Error] Failed workbook {fname}: {wb_err}")
                            summary.append({
                                "source": src_name,
                                "original_table": fname,
                                "sqlite_table": f"reg2025_{prov_key}_*",
                                "rows": 0,
                                "status": f"ERROR: {str(wb_err)[:80]}"
                            })

                elif src_type == "excel":
                    xl = pd.ExcelFile(file_path)
                    for sheet in tables:
                        if sheet not in xl.sheet_names:
                            continue
                        sqlite_tbl_name = f"{prefix}_{sanitize_identifier(sheet)}"
                        _update_sync_state(current_table=f"{src_name} -> {sheet}")
                        t0_tbl = time.time()

                        try:
                            df = pd.read_excel(xl, sheet_name=sheet)
                            rows_copied, sample_fts = ETLPipeline._sync_excel_sheet(
                                df=df,
                                sqlite_conn=sqlite_conn,
                                sqlite_tbl=sqlite_tbl_name,
                                src_name=src_name,
                                sheet_name=sheet
                            )
                            total_rows_copied += rows_copied
                            tables_done += 1
                            fts_entries.extend(sample_fts)
                            _update_sync_state(
                                tables_completed=tables_done,
                                rows_copied=total_rows_copied
                            )
                            elapsed_tbl = time.time() - t0_tbl
                            print(f"[ETL] Extracted {sheet} -> {sqlite_tbl_name}: {rows_copied:,} rows in {elapsed_tbl:.2f}s")
                            summary.append({
                                "source": src_name,
                                "original_table": sheet,
                                "sqlite_table": sqlite_tbl_name,
                                "rows": rows_copied,
                                "status": "SUCCESS"
                            })
                        except Exception as sheet_err:
                            print(f"[ETL Error] Failed sheet {sheet}: {sheet_err}")
                            summary.append({
                                "source": src_name,
                                "original_table": sheet,
                                "sqlite_table": sqlite_tbl_name,
                                "rows": 0,
                                "status": f"ERROR: {str(sheet_err)[:80]}"
                            })

            # Standardize and populate where_farm_taken_from across all tables and create master views
            _update_sync_state(current_table="Standardizing Farm Origins & Lineage...")
            ETLPipeline._post_process_farm_origins(sqlite_conn)

            # Build unified FTS5 search index
            _update_sync_state(current_table="Building Full-Text Search (FTS5) Index...")
            ETLPipeline._build_fts_index(sqlite_conn, fts_entries)

            # Metadata table
            sqlite_conn.execute("CREATE TABLE IF NOT EXISTS _etl_metadata (key TEXT PRIMARY KEY, value TEXT);")
            sqlite_conn.execute("REPLACE INTO _etl_metadata (key, value) VALUES ('last_sync', datetime('now', 'localtime'));")
            sqlite_conn.execute(f"REPLACE INTO _etl_metadata (key, value) VALUES ('total_rows', '{total_rows_copied}');")
            sqlite_conn.commit()

            total_elapsed = time.time() - start_time
            file_size_mb = os.path.getsize(CONSOLIDATED_DB_PATH) / (1024 * 1024)

            _update_sync_state(
                status="completed",
                current_source="Done",
                current_table="",
                tables_completed=tables_done,
                rows_copied=total_rows_copied,
                summary=summary
            )

            result = {
                "success": True,
                "db_path": CONSOLIDATED_DB_PATH,
                "file_size_mb": round(file_size_mb, 2),
                "total_rows": total_rows_copied,
                "tables_completed": tables_done,
                "duration_seconds": round(total_elapsed, 1),
                "summary": summary
            }
            return result

        except Exception as e:
            err_msg = str(e)
            print(f"[ETL Critical Error]: {err_msg}")
            _update_sync_state(status="error", error_message=err_msg)
            return {"success": False, "error": err_msg}
        finally:
            sqlite_conn.close()

    @staticmethod
    def _sync_access_table(cursor, access_tbl: str, sqlite_conn: sqlite3.Connection, sqlite_tbl: str, src_name: str) -> tuple[int, List[tuple]]:
        """Extracts rows in chunks from an Access table and inserts into SQLite"""
        escaped_tbl = f"[{access_tbl}]"
        cursor.execute(f"SELECT TOP 1 * FROM {escaped_tbl}")
        if not cursor.description:
            return 0, []

        raw_col_names = [d[0] for d in cursor.description]
        clean_cols = []
        used_names = set()
        for c in raw_col_names:
            s = sanitize_identifier(c)
            # Avoid collision
            idx = 2
            cand = s
            while cand in used_names:
                cand = f"{s}_{idx}"
                idx += 1
            used_names.add(cand)
            clean_cols.append(cand)

        # Drop and create SQLite table
        sqlite_conn.execute(f"DROP TABLE IF EXISTS [{sqlite_tbl}];")
        col_defs = ", ".join([f"[{col}] TEXT" for col in clean_cols])
        sqlite_conn.execute(f"CREATE TABLE [{sqlite_tbl}] ({col_defs});")

        # Select all rows
        cursor.execute(f"SELECT * FROM {escaped_tbl}")
        placeholders = ", ".join(["?"] * len(clean_cols))
        insert_sql = f"INSERT INTO [{sqlite_tbl}] VALUES ({placeholders})"

        rows_copied = 0
        fts_samples = []
        chunk_size = 1000

        while True:
            rows = cursor.fetchmany(chunk_size)
            if not rows:
                break

            insert_batch = []
            for r in rows:
                row_vals = []
                content_parts = []
                for val in r:
                    if val is None:
                        row_vals.append(None)
                    else:
                        val_str = str(val).strip()
                        row_vals.append(val_str)
                        if val_str:
                            content_parts.append(val_str)
                insert_batch.append(row_vals)

                # Record sample for FTS
                if content_parts:
                    title = str(row_vals[0] or "")
                    subtitle = str(row_vals[1] or "") if len(row_vals) > 1 else ""
                    content = " ".join(content_parts)
                    fts_samples.append((src_name, sqlite_tbl, title, subtitle, content))

            sqlite_conn.executemany(insert_sql, insert_batch)
            rows_copied += len(insert_batch)

        sqlite_conn.commit()
        return rows_copied, fts_samples

    @staticmethod
    def _sync_excel_sheet(df: pd.DataFrame, sqlite_conn: sqlite3.Connection, sqlite_tbl: str, src_name: str, sheet_name: str) -> tuple[int, List[tuple]]:
        """Inserts a pandas DataFrame into SQLite table with FTS sample generation"""
        df.columns = [sanitize_identifier(c) for c in df.columns]
        
        # Deduplicate column names
        cols = []
        used = set()
        for c in df.columns:
            cand = c
            idx = 2
            while cand in used:
                cand = f"{c}_{idx}"
                idx += 1
            used.add(cand)
            cols.append(cand)
        df.columns = cols

        # Clean string representation
        df_str = df.astype(str).replace({'nan': None, 'None': None, '<NA>': None})
        sqlite_conn.execute(f"DROP TABLE IF EXISTS [{sqlite_tbl}];")
        df_str.to_sql(sqlite_tbl, sqlite_conn, index=False, if_exists="replace")

        # Build FTS samples
        fts_samples = []
        for _, row in df_str.iterrows():
            non_empty = [str(v).strip() for v in row if v and str(v).strip() and str(v).strip() != 'None']
            if non_empty:
                title = non_empty[0]
                subtitle = non_empty[1] if len(non_empty) > 1 else ""
                content = " ".join(non_empty)
                fts_samples.append((src_name, sqlite_tbl, title, subtitle, content))

        return len(df), fts_samples

    @staticmethod
    def _sync_register_workbook(file_path: str, prov_key: str, sqlite_conn: sqlite3.Connection, src_name: str) -> tuple[int, List[tuple], int]:
        """Extracts sheets from a 2025 provincial register workbook (.xlsx or .xls) using openpyxl/xlrd"""
        total_rows = 0
        total_tables = 0
        fts_samples = []
        fname = os.path.basename(file_path)

        if fname.endswith('.xlsx') and openpyxl:
            wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
            for sname in wb.sheetnames:
                clean_sname = sanitize_identifier(sname)
                if clean_sname in ['sheet1', 'sheet2', 'stats_by_district']:
                    continue
                tbl_name = f"reg2025_{prov_key}_{clean_sname}"
                ws = wb[sname]
                raw_rows = ws.iter_rows(values_only=True)
                header_row = None
                for r in raw_rows:
                    if any(r):
                        header_row = r
                        break
                if not header_row:
                    continue
                last_idx = 0
                for idx, c in enumerate(header_row):
                    if c is not None and str(c).strip():
                        last_idx = idx
                header_row = header_row[:last_idx + 1]
                cols = []
                used = set()
                for idx, c in enumerate(header_row):
                    cand = sanitize_identifier(c) if c else f"col_{idx+1}"
                    k = 2
                    cand_uniq = cand
                    while cand_uniq in used:
                        cand_uniq = f"{cand}_{k}"
                        k += 1
                    used.add(cand_uniq)
                    cols.append(cand_uniq)
                sqlite_conn.execute(f"DROP TABLE IF EXISTS [{tbl_name}];")
                col_defs = ", ".join([f"[{c}] TEXT" for c in cols])
                sqlite_conn.execute(f"CREATE TABLE [{tbl_name}] ({col_defs});")
                insert_sql = f"INSERT INTO [{tbl_name}] VALUES ({', '.join(['?'] * len(cols))})"
                batch = []
                sheet_rows = 0
                for r in raw_rows:
                    if not any(r):
                        continue
                    vals = [str(v).strip() if v is not None and str(v).strip() != '' and str(v).strip().lower() != 'none' else None for v in r[:len(cols)]]
                    if not any(vals):
                        continue
                    if len(vals) < len(cols):
                        vals.extend([None] * (len(cols) - len(vals)))
                    batch.append(vals)
                    non_empty = [v for v in vals if v]
                    if non_empty:
                        title = non_empty[0]
                        subtitle = non_empty[1] if len(non_empty) > 1 else ""
                        content = " ".join(non_empty[:15])
                        fts_samples.append((src_name, tbl_name, title, subtitle, content))
                    if len(batch) >= 2000:
                        sqlite_conn.executemany(insert_sql, batch)
                        sheet_rows += len(batch)
                        batch = []
                if batch:
                    sqlite_conn.executemany(insert_sql, batch)
                    sheet_rows += len(batch)
                sqlite_conn.commit()
                total_rows += sheet_rows
                total_tables += 1
            wb.close()
        elif fname.endswith('.xls') and xlrd:
            book = xlrd.open_workbook(file_path, on_demand=True)
            for sname in book.sheet_names():
                clean_sname = sanitize_identifier(sname)
                if clean_sname in ['sheet1', 'sheet2', 'stats_by_district']:
                    continue
                tbl_name = f"reg2025_{prov_key}_{clean_sname}"
                s = book.sheet_by_name(sname)
                if s.nrows <= 1:
                    continue
                header_row = [s.cell_value(0, c) for c in range(s.ncols)]
                last_idx = 0
                for idx, c in enumerate(header_row):
                    if c is not None and str(c).strip():
                        last_idx = idx
                header_row = header_row[:last_idx + 1]
                cols = []
                used = set()
                for idx, c in enumerate(header_row):
                    cand = sanitize_identifier(c) if c else f"col_{idx+1}"
                    k = 2
                    cand_uniq = cand
                    while cand_uniq in used:
                        cand_uniq = f"{cand}_{k}"
                        k += 1
                    used.add(cand_uniq)
                    cols.append(cand_uniq)
                sqlite_conn.execute(f"DROP TABLE IF EXISTS [{tbl_name}];")
                col_defs = ", ".join([f"[{c}] TEXT" for c in cols])
                sqlite_conn.execute(f"CREATE TABLE [{tbl_name}] ({col_defs});")
                insert_sql = f"INSERT INTO [{tbl_name}] VALUES ({', '.join(['?'] * len(cols))})"
                batch = []
                sheet_rows = 0
                for row_idx in range(1, s.nrows):
                    row_vals = [s.cell_value(row_idx, c) for c in range(len(cols))]
                    if not any(str(v).strip() for v in row_vals if v is not None):
                        continue
                    vals = [str(v).strip() if v is not None and str(v).strip() != '' and str(v).strip().lower() != 'none' else None for v in row_vals]
                    if not any(vals):
                        continue
                    batch.append(vals)
                    non_empty = [v for v in vals if v]
                    if non_empty:
                        title = non_empty[0]
                        subtitle = non_empty[1] if len(non_empty) > 1 else ""
                        content = " ".join(non_empty[:15])
                        fts_samples.append((src_name, tbl_name, title, subtitle, content))
                    if len(batch) >= 2000:
                        sqlite_conn.executemany(insert_sql, batch)
                        sheet_rows += len(batch)
                        batch = []
                if batch:
                    sqlite_conn.executemany(insert_sql, batch)
                    sheet_rows += len(batch)
                sqlite_conn.commit()
                total_rows += sheet_rows
                total_tables += 1
        return total_rows, fts_samples, total_tables

    @staticmethod
    def _post_process_farm_origins(sqlite_conn: sqlite3.Connection):
        """
        Standardizes 'where_farm_taken_from', 'source_database', and 'source_table'
        across all tables and creates the master view vw_consolidated_farms_origin.
        """
        c = sqlite_conn.cursor()
        tables = [r[0] for r in c.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%' AND name != '_etl_metadata' AND name NOT LIKE 'unified_search_fts%'").fetchall()]

        # 1. Add provenance columns
        for tbl in tables:
            cols = [col[1] for col in c.execute(f"PRAGMA table_info([{tbl}])").fetchall()]
            if tbl.startswith("ol2026_"):
                src_db = "Offer Letters 2026 (Live System ACCDB)"
                src_tbl = tbl[7:]
            elif tbl.startswith("lands02_"):
                src_db = "LANDS 02 Network Database (LIMS Live MDB)"
                src_tbl = tbl[8:]
            elif tbl.startswith("a2_alloc_"):
                src_db = "A2 Allocations Database (Mrs Gavi Excel)"
                src_tbl = tbl[9:]
            elif tbl.startswith("reg2025_"):
                parts = tbl.split("_", 2)
                prov = parts[1].capitalize() if len(parts) > 1 else ""
                sheet = parts[2] if len(parts) > 2 else ""
                src_db = f"Current Farm Registers 2025 ({prov})"
                src_tbl = sheet
            else:
                src_db = "Consolidated Database"
                src_tbl = tbl

            if "where_farm_taken_from" not in cols:
                sqlite_conn.execute(f"ALTER TABLE [{tbl}] ADD COLUMN where_farm_taken_from TEXT;")
            if "source_database" not in cols:
                sqlite_conn.execute(f"ALTER TABLE [{tbl}] ADD COLUMN source_database TEXT;")
            if "source_table" not in cols:
                sqlite_conn.execute(f"ALTER TABLE [{tbl}] ADD COLUMN source_table TEXT;")

            sqlite_conn.execute(f"UPDATE [{tbl}] SET source_database = ?, source_table = ? WHERE source_database IS NULL OR source_database = '';", (src_db, src_tbl))

            # Populate where_farm_taken_from if empty
            has_parent_diag = 'parent_farm_diagm_no' in cols
            has_from_reg = 'from_farm_register' in cols
            has_farm_withdrawn = 'farm_withdrawn' in cols
            has_farm_name = 'farm_name' in cols
            has_farm_name_2 = 'farm_name_2' in cols
            sd_col = 'sd_no' if 'sd_no' in cols else ('subdivision_number' if 'subdivision_number' in cols else ('s_d_no' if 's_d_no' in cols else ('s_d' if 's_d' in cols else None)))
            curr_dist_col = 'current_district' if 'current_district' in cols else ('district' if 'district' in cols else ('disrict' if 'disrict' in cols else None))
            prov_col = 'province' if 'province' in cols else ('unnamed_1' if 'unnamed_1' in cols else None)

            needed_cols = ['rowid']
            for candidate in ['farm_name', 'farm_name_2', 'from_farm_register', 'farm_withdrawn',
                              'parent_farm_diagm_no', 'previous_district', curr_dist_col, prov_col,
                              sd_col, 'diagram_no', 'owner_name', 'company_name', 'deed_no', 'gazette_status']:
                if candidate and candidate in cols and candidate not in needed_cols:
                    needed_cols.append(candidate)

            if len(needed_cols) > 1:
                col_str = ", ".join([f"[{cname}]" for cname in needed_cols])
                rows = c.execute(f"SELECT {col_str} FROM [{tbl}] WHERE where_farm_taken_from IS NULL OR where_farm_taken_from = ''").fetchall()
                updates = []
                for r in rows:
                    d = dict(zip(needed_cols, r))
                    rowid = d['rowid']
                    parts = []

                    p_diag = d.get('parent_farm_diagm_no')
                    if p_diag and str(p_diag).strip() and str(p_diag).lower() not in ['none', 'n/a', '-', '']:
                        parts.append(f"Parent Diagram: {str(p_diag).strip()}")

                    from_reg = d.get('from_farm_register')
                    if from_reg and str(from_reg).strip() and str(from_reg).lower() not in ['from farm register', 'none', 'n/a', '-', '']:
                        parts.append(f"Parent Farm: {str(from_reg).strip()}")

                    fn2 = d.get('farm_name_2')
                    if fn2 and str(fn2).strip() and str(fn2).lower() not in ['none', 'n/a', '-', '']:
                        parts.append(f"Parent Estate: {str(fn2).strip()}")

                    f_withdrawn = d.get('farm_withdrawn')
                    if f_withdrawn and str(f_withdrawn).strip() and str(f_withdrawn).lower() not in ['none', 'n/a', '-', '']:
                        parts.append(f"Withdrawn / Reallocated from: {str(f_withdrawn).strip()}")

                    fn = d.get('farm_name')
                    sd = d.get(sd_col) if sd_col else None
                    if sd and str(sd).strip() and str(sd).lower() not in ['none', 'n/a', '-', 'from subdivision register', '']:
                        if fn and str(fn).strip() and str(fn).lower() not in ['none', 'n/a', '-', '']:
                            parts.append(f"Subdivision {str(sd).strip()} carved out of Farm {str(fn).strip()}")
                        else:
                            parts.append(f"Subdivision {str(sd).strip()}")
                    elif fn and str(fn).strip() and str(fn).lower() not in ['none', 'n/a', '-', 'from farm register', '']:
                        if not any(f"Farm {str(fn).strip()}" in p for p in parts):
                            parts.append(f"Farm: {str(fn).strip()}")

                    prev_d = d.get('previous_district')
                    curr_d = d.get(curr_dist_col) if curr_dist_col else None
                    if prev_d and str(prev_d).strip() and str(prev_d).lower() not in ['none', 'n/a', '-', '']:
                        p_clean = str(prev_d).strip()
                        c_clean = str(curr_d).strip() if curr_d else ""
                        if c_clean and p_clean.lower() != c_clean.lower():
                            parts.append(f"Originally in {p_clean} District (demarcated into {c_clean})")
                        else:
                            parts.append(f"Original District: {p_clean}")
                    elif curr_d and str(curr_d).strip() and str(curr_d).lower() not in ['none', 'n/a', '-', '']:
                        pr_val = d.get(prov_col) if prov_col else None
                        loc = str(curr_d).strip()
                        if pr_val and str(pr_val).strip() and not str(pr_val).strip().isdigit():
                            loc += f", {str(pr_val).strip()}"
                        parts.append(f"Location: {loc}")

                    diag = d.get('diagram_no')
                    if diag and str(diag).strip() and str(diag).lower() not in ['none', 'n/a', '0000', '-', '']:
                        parts.append(f"Diagram: {str(diag).strip()}")

                    deed = d.get('deed_no')
                    if deed and str(deed).strip() and str(deed).lower() not in ['none', 'n/a', '-', '']:
                        parts.append(f"Deed: {str(deed).strip()}")

                    own = d.get('owner_name')
                    if own and str(own).strip() and str(own).lower() not in ['none', 'n/a', '-', 'individual', '']:
                        parts.append(f"Former Owner: {str(own).strip()}")
                    comp = d.get('company_name')
                    if comp and str(comp).strip() and str(comp).lower() not in ['from beneficiary details', 'none', 'n/a', '-', 'individual', '']:
                        parts.append(f"Former Estate/Company: {str(comp).strip()}")

                    gaz = d.get('gazette_status')
                    if gaz and str(gaz).strip() and str(gaz).lower() in ['gazetted']:
                        parts.append("Acquisition: Gazetted by State")

                    where_text = " | ".join(parts) if parts else None
                    if where_text:
                        updates.append((where_text, rowid))

                if updates:
                    sqlite_conn.executemany(f"UPDATE [{tbl}] SET where_farm_taken_from = ? WHERE rowid = ?;", updates)

        # 2. Create master unified farm origin view
        sqlite_conn.execute("DROP VIEW IF EXISTS vw_consolidated_farms_origin;")
        sqlite_conn.execute("""
            CREATE VIEW IF NOT EXISTS vw_consolidated_farms_origin AS
            SELECT 
                'OfferLetters 2026 (Live System)' AS source_database,
                'Farm_register' AS source_table,
                farm_name,
                NULL AS subdivision_number,
                owner_name AS beneficiary_or_owner,
                NULL AS national_id,
                current_district AS district,
                province,
                size_in_hectare AS area_ha,
                where_farm_taken_from
            FROM ol2026_farm_register
            UNION ALL
            SELECT 
                'OfferLetters 2026 (Live System)' AS source_database,
                'PermitsPrinted' AS source_table,
                farm_name,
                subdivision_id AS subdivision_number,
                first_name || ' ' || surname AS beneficiary_or_owner,
                national_id,
                NULL AS district,
                NULL AS province,
                NULL AS area_ha,
                where_farm_taken_from
            FROM ol2026_permitsprinted
            UNION ALL
            SELECT 
                'LANDS 02 Network Database (LIMS)' AS source_database,
                'FARM DETAILS' AS source_table,
                farm_name,
                subdivision_number,
                NULL AS beneficiary_or_owner,
                id_number AS national_id,
                district,
                province,
                subdiv_size AS area_ha,
                where_farm_taken_from
            FROM lands02_farm_details
            UNION ALL
            SELECT 
                'LANDS 02 Network Database (LIMS)' AS source_database,
                'Farm Withdrawn' AS source_table,
                farm_withdrawn AS farm_name,
                subdivision_number,
                NULL AS beneficiary_or_owner,
                id_number AS national_id,
                district,
                province,
                subdiv_size AS area_ha,
                where_farm_taken_from
            FROM lands02_farm_withdrawn
            UNION ALL
            SELECT 
                'A2 Allocations Database (Mrs Gavi)' AS source_database,
                'Sheet1' AS source_table,
                farm_name,
                s_d_no AS subdivision_number,
                name AS beneficiary_or_owner,
                id_number AS national_id,
                disrict AS district,
                unnamed_1 AS province,
                area_ha,
                where_farm_taken_from
            FROM a2_alloc_sheet1
            UNION ALL
            SELECT 
                'A2 Allocations Database (Mrs Gavi)' AS source_database,
                'Sheet2' AS source_table,
                farm_name,
                s_d_no AS subdivision_number,
                surname AS beneficiary_or_owner,
                id_number AS national_id,
                district,
                unnamed_1 AS province,
                area_ha,
                where_farm_taken_from
            FROM a2_alloc_sheet2
            UNION ALL
            SELECT 
                'Current Farm Registers 2025 (Mash West)' AS source_database,
                'A2 Allocation' AS source_table,
                from_farm_register AS farm_name,
                sd_no AS subdivision_number,
                first_name || ' ' || surname AS beneficiary_or_owner,
                id_no AS national_id,
                current_district AS district,
                province,
                extent_ha AS area_ha,
                where_farm_taken_from
            FROM reg2025_mashwest_a2_allocation
            UNION ALL
            SELECT 
                'Current Farm Registers 2025 (Masvingo)' AS source_database,
                'A1 Allocation' AS source_table,
                farm_name,
                s_d AS subdivision_number,
                first_name || ' ' || surname AS beneficiary_or_owner,
                idnumber AS national_id,
                current_district AS district,
                province,
                NULL AS area_ha,
                where_farm_taken_from
            FROM reg2025_masvingo_a1_allocation
            UNION ALL
            SELECT 
                'Current Farm Registers 2025 (Mash East)' AS source_database,
                'A1 Allocation' AS source_table,
                farm_name,
                sd_no AS subdivision_number,
                first_nme || ' ' || surname AS beneficiary_or_owner,
                id_number AS national_id,
                current_district AS district,
                province,
                extent_ha AS area_ha,
                where_farm_taken_from
            FROM reg2025_masheast_a1_allocation
            UNION ALL
            SELECT 
                'Current Farm Registers 2025 (Midlands)' AS source_database,
                'A1 Allocation' AS source_table,
                farm_name,
                sd_no AS subdivision_number,
                NULL AS beneficiary_or_owner,
                NULL AS national_id,
                current_district AS district,
                province,
                NULL AS area_ha,
                where_farm_taken_from
            FROM reg2025_midlands_a1_allocation;
        """)
        sqlite_conn.commit()

    @staticmethod
    def _build_fts_index(sqlite_conn: sqlite3.Connection, fts_entries: List[tuple]):
        """Creates and populates the SQLite FTS5 Full-Text Search virtual table"""
        sqlite_conn.execute("DROP TABLE IF EXISTS unified_search_fts;")
        sqlite_conn.execute("""
            CREATE VIRTUAL TABLE unified_search_fts USING fts5(
                source_name,
                table_name,
                title,
                subtitle,
                content
            );
        """)

        # Insert in chunks
        chunk_size = 2000
        for i in range(0, len(fts_entries), chunk_size):
            chunk = fts_entries[i:i + chunk_size]
            sqlite_conn.executemany(
                "INSERT INTO unified_search_fts (source_name, table_name, title, subtitle, content) VALUES (?, ?, ?, ?, ?)",
                chunk
            )
        sqlite_conn.commit()

def start_sync_background() -> bool:
    """Spawns ETL sync in a background daemon thread if not already running"""
    with _SYNC_LOCK:
        if _SYNC_STATE["status"] == "running":
            return False

    t = threading.Thread(target=ETLPipeline.run_full_sync, daemon=True)
    t.start()
    return True

if __name__ == "__main__":
    print(f"Starting ETL pipeline consolidation into {CONSOLIDATED_DB_PATH}...")
    res = ETLPipeline.run_full_sync()
    print("\nResult:", res)
